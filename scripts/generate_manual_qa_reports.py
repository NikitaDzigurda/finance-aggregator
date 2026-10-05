"""Reproducible, fictional provider-shaped reports. Never reads personal reports."""

from __future__ import annotations

import csv
import json
from collections import defaultdict
from copy import deepcopy
from decimal import Decimal as D
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / "docs/fixtures/manual_qa"


def broker_trades(prefix: str):
    # Five instruments, each bought twice and partially sold. All start at zero.
    rows = []
    for i in range(1, 6):
        code = f"RU000QA{prefix}{i:04d}"
        name = f"QA {prefix} Акция {i} — синтетический пример"
        for month, quantity, price in ((1, 10, 100 * i), (2, 5, 110 * i), (3, -3, 120 * i)):
            rows.append(
                dict(
                    code=code,
                    name=name,
                    date=f"2026-{month:02d}-{i + 5:02d}",
                    qty=D(quantity),
                    price=D(price),
                    fee=D("1.25"),
                    id=f"QA-{prefix}-{month}-{i}",
                )
            )
    return sorted(rows, key=lambda r: r["date"])


def broker_expected(trades, income=D(0)):
    cash = D(100000) + income
    quantities = defaultdict(D)
    for t in trades:
        cash -= t["qty"] * t["price"] + t["fee"]
        quantities[t["code"]] += t["qty"]
    return {
        "cash_RUB": str(cash),
        "quantities": {k: str(v) for k, v in quantities.items()},
        "fees_RUB": str(sum(t["fee"] for t in trades)),
        "open_cost_basis_RUB_excluding_fees": "18600",
        "realized_pnl_RUB_excluding_fees": "750",
    }


def tbank():
    trades = broker_trades("T")
    expected = broker_expected(trades)
    wb = Workbook()
    wb.remove(wb.active)
    for name in (
        "Динамика позиций",
        "Завершенные сделки",
        "Незавершенные сделки",
        " Движение ДС",
        "Неторговые операции",
    ):
        wb.create_sheet(name)
    ws = wb["Завершенные сделки"]
    ws.append(["СИНТЕТИЧЕСКИЙ ОТЧЁТ — не данные клиента"])
    ws.append(
        [
            "№ сделки",
            "Дата заключен.",
            "ISIN/рег.код",
            "Актив",
            "Количество актива⁷, шт./грамм",
            "Цена",
            "Валюта расчетов",
            "Комиссия банка",
            "Валюта комиссии",
        ]
    )
    for t in trades:
        year, month, day = t["date"].split("-")
        ws.append(
            [
                t["id"],
                f"{day}.{month}.{year} 12:00:00",
                t["code"],
                t["name"],
                t["qty"],
                t["price"],
                "RUB",
                t["fee"],
                "RUB",
            ]
        )
    ws = wb["Динамика позиций"]
    ws.append(["Инструмент", "Актив ¹", "кол-во, шт./грамм⁵"])
    ws.append(["Входящий остаток / Исходящий остаток"])
    for t in trades[:5]:
        row = [None] * 16
        row[2], row[6], row[12], row[15] = "Акции", t["name"], 0, 12
        ws.append(row)
    ws = wb[" Движение ДС"]
    ws.cell(1, 15, 0)
    ws.cell(2, 1, "RUB")
    ws.append(["Дата", "Наименование операции", "Комментарий"])
    ws.append(["Движение денежных средств"])
    row = [None] * 18
    row[6], row[9], row[10], row[14] = "02.01.2026", "Перевод", "QA пополнение", D(100000)
    ws.append(row)
    # Settlement lines are informational in this layout, not additional deposits.
    cash = D(100000)
    for t in trades:
        cash -= t["qty"] * t["price"] + t["fee"]
        row = [None] * 18
        row[6], row[9], row[10] = t["date"], "Расчеты по сделке", t["id"]
        row[14], row[17] = -(t["qty"] * t["price"] + t["fee"]), cash
        ws.append(row)
    for ws in wb:
        ws.freeze_panes = "A3"
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="153D77")
        for column in "ABCDEFGHIJKLMNO PQR".replace(" ", ""):
            ws.column_dimensions[column].width = 24
    wb.save(DEST / "Т-Инвестиции-SYNTHETIC-QA-(01.01.26-31.03.26).xlsx")
    return expected


def set_fields(node, **values):
    for key, value in values.items():
        child = node.find(key)
        if child is None:
            child = ET.SubElement(node, key)
        child.text = str(value)


def alfa():
    root = ET.parse(ROOT / "docs/fixtures/alfa_broker_report_import_synthetic_v1.xml").getroot()
    trade_template = deepcopy(root.find("trades_finished/trade"))
    position_template = deepcopy(root.find("positions/position"))
    money_template = deepcopy(root.find("money_moves/money_move"))
    for section in (
        "trades_finished",
        "trades_unfinished",
        "positions",
        "money_moves",
        "transfers",
    ):
        root.find(section).clear()
    set_fields(
        root,
        full_name="SYNTHETIC QA",
        treaty="SYNTHETIC-QA-ALFA",
        date_start="2026-01-01",
        date_end="2026-03-31",
    )
    trades = broker_trades("A")
    expected = broker_expected(trades, D(250))
    totals = root.find("money_moves_total")
    totals.clear()
    total_values = dict(
        begin_real_rest="0",
        end_real_rest=expected["cash_RUB"],
        saldo_rest=expected["cash_RUB"],
        extra_transfer_in="100000",
        extra_transfer_out="0",
        dividend="250",
        coupon="0",
        comission=expected["fees_RUB"],
        trade_transfer_in=str(sum(-t["qty"] * t["price"] for t in trades if t["qty"] < 0)),
        trade_transfer_out=str(sum(t["qty"] * t["price"] for t in trades if t["qty"] > 0)),
    )
    for tag, value in total_values.items():
        child = ET.SubElement(ET.SubElement(totals, tag + "_total"), tag)
        set_fields(child, p_code="RUB", value=value)
    cash = D(0)

    def money(date, amount, group, comment, trade_no=None):
        nonlocal cash
        node = deepcopy(money_template)
        set_fields(
            node,
            settlement_date=date + "T12:00:00",
            volume=amount,
            oper_group=group,
            comment=comment,
            begin_real_rest=cash,
            money_volume_begin=cash,
            all_volume_begin=cash,
        )
        cash += amount
        set_fields(node, end_real_rest=cash, money_volume_end=cash, all_volume_end=cash)
        if trade_no:
            set_fields(node, trd_no=trade_no, oper_type="Расчеты по сделке")
        root.find("money_moves").append(node)

    money("2026-01-02", D(100000), "Внесено по распоряжению Клиента", "QA пополнение")
    for t in trades:
        node = deepcopy(trade_template)
        set_fields(
            node,
            db_time=t["date"] + "T12:00:00",
            settlement_time=t["date"] + "T12:00:00",
            p_name=t["name"],
            Price=t["price"],
            summ_trade=abs(t["qty"] * t["price"]),
            qty=t["qty"],
            bank_tax=t["fee"],
            isin_reg=t["code"],
            trade_no=t["id"],
        )
        root.find("trades_finished").append(node)
        money(t["date"], -t["qty"] * t["price"], "", "QA расчёт", t["id"])
        money(t["date"], -t["fee"], "", "QA комиссия", t["id"])
    money("2026-03-25", D(250), "Дивиденды", "QA выплата")
    for i, t in enumerate(trades[:5], 1):
        node = deepcopy(position_template)
        set_fields(
            node,
            active_rank=i,
            active_type="Акция",
            active_name=t["name"],
            ISIN=t["code"],
            p_code=f"QAA{i}",
            income_rest=0,
            real_rest=12,
            forward_rest=12,
            real_volume=D(12) * D(120 * i),
            at_group="Акции",
        )
        set_fields(node, **{"Цена_на_начало_периода": "0", "Цена_на_конец_периода": str(120 * i)})
        root.find("positions").append(node)
    assert cash == D(expected["cash_RUB"])
    ET.indent(root)
    ET.ElementTree(root).write(
        DEST / "Альфа-SYNTHETIC-QA-(01.01.26-31.03.26).xml", encoding="utf-8", xml_declaration=True
    )
    return expected


def bybit():
    names = {
        "spot": ("bybit_spot_trade_history_synthetic_v1.csv", "unifiedAccount_spotTradeHistory"),
        "uta": ("bybit_uta_asset_change_details_synthetic_v1.csv", "AssetChangeDetails_uta"),
        "fund": ("bybit_funding_asset_change_details_synthetic_v1.csv", "AssetChangeDetails_fund"),
        "chain": (
            "bybit_withdraw_deposit_history_synthetic_v1.csv",
            "assetHistory_withdrawDepositHistory",
        ),
    }
    rows = {key: [] for key in names}
    uid = "900000099"
    balances = defaultdict(D, USDT=D(4000))
    rows["fund"] = [
        [uid, "2026-01-02 09:00:00", "USDT", "5000", "Fiat", "5000", "P2P Purchase"],
        [
            uid,
            "2026-01-03 09:00:00",
            "USDT",
            "-4000",
            "Transfer out",
            "1000",
            "Transfer to Unified Trading Account",
        ],
    ]
    rows["uta"].append(
        [
            uid,
            "USDT",
            "",
            "TRANSFER_IN",
            "--",
            "0",
            "0",
            "0",
            "0",
            "0",
            "4000",
            "4000",
            "4000",
            "--",
            "2026-01-03 09:00:00",
        ]
    )
    for n in range(12):
        coin = "QAX" if n % 2 == 0 else "QAY"
        side = "SELL" if n >= 8 else "BUY"
        quantity = D(2 if side == "SELL" else 5)
        price = D(20 + n)
        value = quantity * price
        fee = quantity * D("0.001") if side == "BUY" else value * D("0.001")
        date = f"2026-{1 + n // 4:02d}-{10 + n % 4:02d} 10:00:00"
        rows["spot"].append(
            [
                uid,
                coin + "USDT",
                "MARKET",
                side,
                value,
                price,
                quantity,
                fee,
                str(99000000 + n),
                str(99100000 + n),
                date,
            ]
        )
        for currency, amount, paid in (
            (coin, quantity if side == "BUY" else -quantity, -fee if side == "BUY" else D(0)),
            ("USDT", -value if side == "BUY" else value, -fee if side == "SELL" else D(0)),
        ):
            change = amount + paid
            balances[currency] += change
            rows["uta"].append(
                [
                    uid,
                    currency,
                    coin + "USDT",
                    "TRADE",
                    side,
                    amount,
                    "0",
                    price,
                    "0",
                    paid,
                    amount,
                    change,
                    balances[currency],
                    "--",
                    date,
                ]
            )
    rows["fund"].append(
        [uid, "2026-03-20 12:00:00", "USDT", "-100.1", "Withdraw", "899.9", "Withdrawal"]
    )
    rows["chain"].append(
        [
            uid,
            "2026-03-20 11:59:59",
            "Withdraw",
            "USDT",
            "SYNTH_CHAIN",
            "100",
            "0x" + "c" * 64,
            "Transferred successfully",
            "0x" + "d" * 40,
        ]
    )
    for key, (template, prefix) in names.items():
        with (ROOT / "docs/fixtures" / template).open(encoding="utf-8-sig") as f:
            lines = list(csv.reader(f))
        with (DEST / f"{prefix}_{uid}_20260101_20260331_0.csv").open(
            "w", encoding="utf-8", newline=""
        ) as f:
            writer = csv.writer(f)
            writer.writerow([f"UID: {uid}", "Company Name: SYNTHETIC QA", "Country: ZZ"])
            writer.writerow(lines[1])
            writer.writerows(rows[key])
    balances["USDT"] += D("899.9")
    return {
        "quantities": {k: str(v) for k, v in balances.items()},
        "trade_count": 12,
        "cost_basis": "unknown: fiat purchase has no fiat amount in these CSVs",
    }


def main():
    DEST.mkdir(parents=True, exist_ok=True)
    expected = {"tbank": tbank(), "alfa": alfa(), "bybit": bybit()}
    (DEST / "expected.json").write_text(
        json.dumps(expected, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    output = ROOT / "outputs/manual-qa-reports-v2.zip"
    with ZipFile(output, "w", ZIP_DEFLATED) as z:
        for path in sorted(DEST.iterdir()):
            if path.is_file():
                z.write(path, path.name)
    print(output)


if __name__ == "__main__":
    main()
