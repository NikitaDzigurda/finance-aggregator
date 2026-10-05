import os
from collections.abc import AsyncIterator
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook
from sqlalchemy import delete, func, select, update
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from apps.api.main import create_app
from imports.adapters import get_adapter_registry
from imports.jobs import claim_next_job, finish_job
from imports.models import ImportBatchModel, ImportRowModel, ImportRowStatus
from imports.processor import ImportProcessingError, fail_import_batch, process_import_batch
from imports.storage import LocalObjectStorage, get_object_storage
from instruments.models import InstrumentModel
from operations.models import OperationModel
from portfolios.models import PortfolioModel
from shared.database import get_db_session

_FIXTURE = (
    Path(__file__).parents[1]
    / "docs"
    / "fixtures"
    / "tbank_broker_report_synthetic_v1.xlsx"
)


async def _upload_and_process(
    client: AsyncClient,
    sessions: async_sessionmaker[AsyncSession],
    storage: LocalObjectStorage,
    *,
    portfolio_id: UUID,
    account_id: str,
    content: bytes,
    filename: str,
) -> UUID:
    response = await client.post(
        "/api/v1/imports",
        data={
            "portfolio_id": str(portfolio_id),
            "account_id": account_id,
            "source_provider": "tbank_broker_xlsx",
            "declared_format": "xlsx",
        },
        files={
            "file": (
                filename,
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert response.status_code == 201, response.text
    batch_id = UUID(response.json()["batch"]["id"])
    worker_id = f"tbank-test-{uuid4()}"
    async with sessions() as session:
        job = await claim_next_job(session, worker_id=worker_id)
    assert job is not None
    assert job.batch_id == batch_id
    await process_import_batch(
        sessions,
        storage,
        get_adapter_registry(),
        batch_id=batch_id,
    )
    async with sessions() as session:
        await finish_job(
            session,
            job_id=job.id,
            worker_id=worker_id,
            succeeded=True,
        )
    return batch_id


@pytest.mark.postgres
@pytest.mark.asyncio
async def test_tbank_full_import_overlap_recalculate_and_rollback(tmp_path: Path) -> None:
    database_url = os.getenv("FINANCE_TEST_DATABASE_URL")
    if database_url is None:
        pytest.skip("FINANCE_TEST_DATABASE_URL is not configured")

    engine = create_async_engine(database_url)
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    storage = LocalObjectStorage(tmp_path / "tbank-imports", max_file_size_bytes=128 * 1024)

    async def override_session() -> AsyncIterator[AsyncSession]:
        async with sessions() as session:
            try:
                yield session
            except Exception:
                await session.rollback()
                raise

    application = create_app()
    application.dependency_overrides[get_db_session] = override_session
    application.dependency_overrides[get_object_storage] = lambda: storage
    transport = ASGITransport(app=application)
    portfolio_id: UUID | None = None
    instrument_id: UUID | None = None
    batch_ids: list[UUID] = []
    storage_keys: list[str] = []

    try:
        async with AsyncClient(transport=transport, base_url="http://test") as client:
            portfolio_response = await client.post(
                "/api/v1/portfolios",
                json={"name": f"Tbank {uuid4().hex[:8]}", "base_currency": "RUB"},
            )
            assert portfolio_response.status_code == 201
            portfolio_id = UUID(portfolio_response.json()["id"])
            account_response = await client.post(
                f"/api/v1/portfolios/{portfolio_id}/accounts",
                json={"name": "Synthetic T broker", "account_type": "broker"},
            )
            assert account_response.status_code == 201
            account_id = account_response.json()["id"]
            instrument_response = await client.post(
                "/api/v1/instruments",
                json={
                    "name": "Synthetic Bond Fund",
                    "instrument_type": "fund",
                    "currency": "RUB",
                    "identifiers": [
                        {
                            "identifier_type": "provider_code",
                            "value": "SYNBOND",
                            "provider": "tbank_broker_xlsx",
                        }
                    ],
                },
            )
            assert instrument_response.status_code == 201, instrument_response.text
            instrument_id = UUID(instrument_response.json()["id"])

            first_batch = await _upload_and_process(
                client,
                sessions,
                storage,
                portfolio_id=portfolio_id,
                account_id=account_id,
                content=_FIXTURE.read_bytes(),
                filename="synthetic-tbank.xlsx",
            )
            batch_ids.append(first_batch)
            status_response = await client.get(f"/api/v1/imports/{first_batch}")
            status_body = status_response.json()["batch"]
            assert status_body["status"] == "ready_to_commit"
            assert status_body["reporting_period_start"] == "2026-01-01"
            assert status_body["reporting_period_end"] == "2026-01-31"
            assert status_body["completeness"] == "period_ledger"
            assert status_body["reconciliation_status"] == "mismatch"

            preview = (await client.get(f"/api/v1/imports/{first_batch}/preview")).json()
            assert preview["total_rows"] == 4
            assert preview["excluded_rows"] == 2
            assert preview["summary"]["operation_counts"] == {"fee": 1, "trade": 1}
            assert preview["reconciliation_summary"] == {
                "status": "mismatch",
                "check_count": 2,
                "mismatch_count": 1,
                "diagnostics": [
                    {
                        "code": "import_reconciliation_cash_mismatch",
                        "message": "Calculated cash movement does not match report controls",
                        "count": 1,
                    }
                ],
            }
            assert {
                row["normalized_candidate"]["payload"]["instrument_id"]
                for row in preview["items"]
                if row["normalized_candidate"] is not None
            } == {str(instrument_id)}

            confirm = await client.post(f"/api/v1/imports/{first_batch}/confirm")
            assert confirm.status_code == 200, confirm.text
            assert confirm.json()["operation_count"] == 2
            assert confirm.json()["idempotent"] is False
            repeated_confirm = await client.post(f"/api/v1/imports/{first_batch}/confirm")
            assert repeated_confirm.json()["idempotent"] is True

            recalculation = await client.post(
                f"/api/v1/portfolios/{portfolio_id}/positions/recalculate",
                json={
                    "as_of": "2026-02-01T00:00:00Z",
                    "cost_basis_method": "weighted_average",
                },
            )
            assert recalculation.status_code == 200, recalculation.text
            assert Decimal(recalculation.json()["positions"][0]["quantity"]) == Decimal("10")

            overview = await client.get(
                f"/api/v1/portfolios/{portfolio_id}/analytics/overview",
                params={"reporting_currency": "RUB"},
            )
            holdings = await client.get(
                f"/api/v1/portfolios/{portfolio_id}/analytics/holdings",
                params={"reporting_currency": "RUB"},
            )
            quality = await client.get(
                f"/api/v1/portfolios/{portfolio_id}/analytics/data-quality",
                params={"reporting_currency": "RUB"},
            )
            assert overview.status_code == 200, overview.text
            assert holdings.status_code == 200, holdings.text
            assert quality.status_code == 200, quality.text
            assert overview.json()["current_value"]["quality"] == "partial"
            assert any(
                diagnostic["code"] == "market_price_missing"
                for item in holdings.json()["items"]
                for diagnostic in item["diagnostics"]
            )
            quality_by_code = {
                item["code"]: item for item in quality.json()["diagnostics"]
            }
            assert quality_by_code["import_history_period_limited"]["count"] == 1
            reconciliation_quality = quality_by_code["import_reconciliation_mismatch"]
            assert reconciliation_quality["count"] == 1
            assert set(reconciliation_quality) == {
                "severity",
                "code",
                "message",
                "count",
                "impacts",
            }
            missing_price_quality = quality_by_code["market_price_missing"]
            missing_price_impacts = {
                (item["endpoint"], item["metric"]): item["state"]
                for item in missing_price_quality["impacts"]
            }
            assert missing_price_impacts[("analytics/overview", "current_value")] == (
                "partial"
            )
            assert missing_price_impacts[("analytics/overview", "cost_basis")] == (
                "complete"
            )
            safe_quality_text = quality.text
            assert "reconciliation_summary" not in safe_quality_text
            assert "instrument_id" not in safe_quality_text
            assert "account_id" not in safe_quality_text
            assert "raw_data" not in safe_quality_text

            workbook = load_workbook(_FIXTURE)
            workbook["broker_rep"]["A1"] = "SYNTHETIC OVERLAP FIXTURE"
            sheet = workbook["broker_rep"]
            sheet.insert_rows(11)
            for column in range(1, sheet.max_column + 1):
                sheet.cell(11, column).value = sheet.cell(10, column).value
                if sheet.cell(9, column).value == "Номер сделки":
                    sheet.cell(11, column).value = "SYN-TRADE-0002"
            overlap_stream = BytesIO()
            workbook.save(overlap_stream)
            workbook.close()
            overlap_batch = await _upload_and_process(
                client,
                sessions,
                storage,
                portfolio_id=portfolio_id,
                account_id=account_id,
                content=overlap_stream.getvalue(),
                filename="synthetic-tbank-overlap.xlsx",
            )
            batch_ids.append(overlap_batch)
            overlap_status = await client.get(f"/api/v1/imports/{overlap_batch}")
            assert overlap_status.json()["batch"]["duplicate_rows"] == 2
            assert overlap_status.json()["batch"]["ready_rows"] == 2
            rejected_confirm = await client.post(f"/api/v1/imports/{overlap_batch}/confirm")
            assert rejected_confirm.status_code == 409

            unresolved_quality = await client.get(
                f"/api/v1/portfolios/{portfolio_id}/analytics/data-quality",
                params={"reporting_currency": "RUB"},
            )
            assert unresolved_quality.status_code == 200, unresolved_quality.text
            unresolved_by_code = {
                item["code"]: item
                for item in unresolved_quality.json()["diagnostics"]
            }
            assert unresolved_by_code["import_review_unresolved"]["count"] == 2

            overlap_preview = (
                await client.get(f"/api/v1/imports/{overlap_batch}/preview")
            ).json()
            duplicate_rows = [
                row for row in overlap_preview["items"] if row["status"] == "duplicate"
            ]
            assert {row["normalized_candidate"]["operation_type"] for row in duplicate_rows} == {
                "trade",
                "fee",
            }
            for duplicate_row in duplicate_rows:
                excluded = await client.patch(
                    f"/api/v1/imports/{overlap_batch}/rows/{duplicate_row['id']}",
                    json={"action": "exclude", "note": "Already imported overlap row"},
                )
                assert excluded.status_code == 200, excluded.text

            async with sessions.begin() as session:
                control_row = await session.scalar(
                    select(ImportRowModel).where(
                        ImportRowModel.batch_id == overlap_batch,
                        ImportRowModel.status == ImportRowStatus.EXCLUDED,
                        ImportRowModel.normalized_candidate.is_(None),
                        ImportRowModel.reconciliation_data.is_not(None),
                    )
                )
                assert control_row is not None
                original_reconciliation_data = control_row.reconciliation_data
                control_row_id = control_row.id
                control_row.reconciliation_data = {
                    "controls": [
                        {
                            "kind": "cash",
                            "key": "RUB",
                            "opening": "NaN",
                            "closing": "1",
                        }
                    ]
                }

            invalid_reconciliation = await client.post(
                f"/api/v1/imports/{overlap_batch}/confirm"
            )
            assert invalid_reconciliation.status_code == 409
            assert invalid_reconciliation.json()["error"]["code"] == (
                "import_reconciliation_invalid"
            )

            async with sessions() as session:
                operation_count = await session.scalar(
                    select(func.count()).select_from(OperationModel).where(
                        OperationModel.account_id == UUID(account_id)
                    )
                )
            assert operation_count == 2

            async with sessions.begin() as session:
                control_row = await session.get(ImportRowModel, control_row_id)
                assert control_row is not None
                control_row.reconciliation_data = original_reconciliation_data

            overlap_confirm = await client.post(f"/api/v1/imports/{overlap_batch}/confirm")
            assert overlap_confirm.status_code == 200, overlap_confirm.text
            assert overlap_confirm.json()["operation_count"] == 2
            assert overlap_confirm.json()["idempotent"] is False

            async with sessions() as session:
                operation_count = await session.scalar(
                    select(func.count()).select_from(OperationModel).where(
                        OperationModel.account_id == UUID(account_id)
                    )
                )
            assert operation_count == 4

            ledger_trades = await client.get(
                "/api/v1/operations",
                params={
                    "portfolio_id": str(portfolio_id),
                    "operation_type": "trade",
                    "limit": 100,
                },
            )
            assert ledger_trades.status_code == 200, ledger_trades.text
            expected_turnover = sum(
                (
                    Decimal(item["payload"]["quantity"])
                    * Decimal(item["payload"]["price"])
                    for item in ledger_trades.json()["items"]
                ),
                start=Decimal(0),
            )
            overlap_timeline = await client.get(
                f"/api/v1/portfolios/{portfolio_id}/analytics/trading",
                params={
                    "from": "2026-01-01",
                    "to": "2026-02-02",
                    "bucket": "month",
                    "timezone": "Europe/Moscow",
                },
            )
            assert overlap_timeline.status_code == 200, overlap_timeline.text
            turnover_series = next(
                item
                for item in overlap_timeline.json()["series"]
                if item["metric"] == "trade_turnover"
            )
            assert sum(
                (Decimal(item["value"]) for item in turnover_series["buckets"]),
                start=Decimal(0),
            ) == expected_turnover

            wrong_provider_workbook = load_workbook(_FIXTURE)
            wrong_provider_workbook["broker_rep"]["A1"] = (
                "SYNTHETIC WRONG PROVIDER FIXTURE"
            )
            wrong_provider_stream = BytesIO()
            wrong_provider_workbook.save(wrong_provider_stream)
            wrong_provider_workbook.close()
            wrong_provider_upload = await client.post(
                "/api/v1/imports",
                data={
                    "portfolio_id": str(portfolio_id),
                    "account_id": account_id,
                    "source_provider": "alfa_broker_xml_import",
                    "declared_format": "xlsx",
                },
                files={
                    "file": (
                        "synthetic-wrong-provider.xlsx",
                        wrong_provider_stream.getvalue(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert wrong_provider_upload.status_code == 201, wrong_provider_upload.text
            wrong_provider_batch = UUID(wrong_provider_upload.json()["batch"]["id"])
            batch_ids.append(wrong_provider_batch)
            worker_id = f"wrong-provider-{uuid4()}"
            async with sessions() as session:
                job = await claim_next_job(session, worker_id=worker_id)
            assert job is not None
            assert job.batch_id == wrong_provider_batch
            with pytest.raises(ImportProcessingError) as captured:
                await process_import_batch(
                    sessions,
                    storage,
                    get_adapter_registry(),
                    batch_id=wrong_provider_batch,
                )
            assert captured.value.code == "import_adapter_not_found"
            await fail_import_batch(
                sessions,
                batch_id=wrong_provider_batch,
                code=captured.value.code,
                message=captured.value.message,
            )
            async with sessions() as session:
                await finish_job(
                    session,
                    job_id=job.id,
                    worker_id=worker_id,
                    succeeded=False,
                    error={"code": captured.value.code, "message": captured.value.message},
                )
            wrong_provider_status = await client.get(
                f"/api/v1/imports/{wrong_provider_batch}"
            )
            assert wrong_provider_status.json()["batch"]["status"] == "failed"
            assert wrong_provider_status.json()["batch"]["error_summary"]["code"] == (
                "import_adapter_not_found"
            )

            rollback = await client.post(f"/api/v1/imports/{first_batch}/rollback")
            assert rollback.status_code == 200, rollback.text
            assert rollback.json()["rolled_back_operations"] == 2
            assert rollback.json()["idempotent"] is False
            repeated_rollback = await client.post(f"/api/v1/imports/{first_batch}/rollback")
            assert repeated_rollback.json()["idempotent"] is True
    finally:
        async with sessions.begin() as session:
            if portfolio_id is not None:
                scoped_batches = select(ImportBatchModel.id).where(
                    ImportBatchModel.portfolio_id == portfolio_id
                )
                storage_keys = list(
                    await session.scalars(
                        select(ImportBatchModel.storage_key).where(
                            ImportBatchModel.portfolio_id == portfolio_id
                        )
                    )
                )
                await session.execute(
                    update(ImportRowModel)
                    .where(ImportRowModel.batch_id.in_(scoped_batches))
                    .values(matched_operation_id=None)
                )
                await session.execute(
                    delete(OperationModel).where(
                        OperationModel.import_batch_id.in_(scoped_batches)
                    )
                )
                await session.execute(
                    delete(ImportBatchModel).where(
                        ImportBatchModel.portfolio_id == portfolio_id
                    )
                )
                await session.execute(
                    delete(PortfolioModel).where(PortfolioModel.id == portfolio_id)
                )
            if instrument_id is not None:
                await session.execute(
                    delete(InstrumentModel).where(InstrumentModel.id == instrument_id)
                )
        for storage_key in storage_keys:
            storage.delete(storage_key)
        application.dependency_overrides.clear()
        await engine.dispose()
