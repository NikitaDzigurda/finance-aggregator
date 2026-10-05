from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile, ZipFile
from zoneinfo import ZoneInfo

from openpyxl import load_workbook  # type: ignore[import-untyped]
from pydantic import TypeAdapter, ValidationError

from imports.adapters import (
    DetectionResult,
    ImportDocument,
    ParsedImport,
    ParsedRow,
    ValidationResult,
)
from imports.models import ImportCompleteness, ImportFileFormat, ImportRowStatus
from shared.exact import CurrencyCode, Money, Price, Quantity, decimal_to_json

type Diagnostic = dict[str, object]

_REPORT_TITLE = re.compile(
    r"^Отчет о сделках и операциях за период "
    r"(?P<start>\d{2}\.\d{2}\.\d{4}) - (?P<end>\d{2}\.\d{2}\.\d{4})$"
)
_MULTISHEET_PERIOD = re.compile(
    r"\((?P<start>\d{2}\.\d{2}\.\d{2,4})-(?P<end>\d{2}\.\d{2}\.\d{2,4})\)"
)
_ISIN = re.compile(r"^[A-Z]{2}[A-Z0-9]{9}[0-9]$")
_TIMEZONE = ZoneInfo("Europe/Moscow")
_MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
_MAX_ARCHIVE_ENTRIES = 2048
_MULTISHEET_REQUIRED_SHEETS = frozenset(
    {
        "Динамика позиций",
        "Завершенные сделки",
        "Незавершенные сделки",
        " Движение ДС",
        "Неторговые операции",
    }
)
_REQUIRED_TRADE_HEADERS = frozenset(
    {
        "Номер сделки",
        "Дата заключения",
        "Время",
        "Вид сделки",
        "Наименование актива",
        "Код актива",
        "Цена за единицу",
        "Валюта цены",
        "Количество",
        "Комиссия брокера",
        "Валюта комиссии",
    }
)
_FEE_FIELDS = (
    ("Комиссия брокера", "Валюта комиссии"),
    ("Комиссия биржи", "Валюта комиссии биржи"),
    ("Комиссия клир. центра", "Валюта комиссии клир. центра"),
)
_KNOWN_SECTION_PREFIXES = (
    "1.1 Информация о совершенных и исполненных сделках",
    "1.2 Информация о неисполненных сделках",
    "1.3 Сделки за расчетный период",
    "2. Операции с денежными средствами",
    "3.1 Движение по ценным бумагам инвестора",
    "3.2 Движение по производным",
    "3.4 Информация по начисленной",
    "4.1 Информация о ценных бумагах",
)
_MONEY_ADAPTER: TypeAdapter[Decimal] = TypeAdapter(Money)
_QUANTITY_ADAPTER: TypeAdapter[Decimal] = TypeAdapter(Quantity)
_PRICE_ADAPTER: TypeAdapter[Decimal] = TypeAdapter(Price)
_CURRENCY_ADAPTER: TypeAdapter[str] = TypeAdapter(CurrencyCode)


@dataclass(frozen=True, slots=True)
class _SourceRow:
    source_row_number: int
    values: dict[str, object]
    source_sheet: str = "broker_rep"


@dataclass(frozen=True, slots=True)
class _InstrumentMetadata:
    name: str
    instrument_type: str


@dataclass(frozen=True, slots=True)
class _WorkbookLayout:
    matched: bool
    reporting_period_start: date | None = None
    reporting_period_end: date | None = None
    trade_rows: tuple[_SourceRow, ...] = ()
    cash_movement_rows: tuple[_SourceRow, ...] = ()
    cash_rows: tuple[_SourceRow, ...] = ()
    security_rows: tuple[_SourceRow, ...] = ()
    instrument_isins: tuple[tuple[str, str], ...] = ()
    instrument_metadata: tuple[tuple[str, _InstrumentMetadata], ...] = ()
    diagnostics: tuple[Diagnostic, ...] = ()


class TbankBrokerXlsxAdapter:
    format_id = "tbank_broker_xlsx"
    version = "1.0"
    supported_file_formats = frozenset({ImportFileFormat.XLSX})

    def detect(self, document: ImportDocument) -> DetectionResult:
        layout = _read_layout(document)
        return DetectionResult(
            matched=layout.matched,
            completeness=(
                ImportCompleteness.PERIOD_LEDGER
                if layout.matched
                else ImportCompleteness.UNKNOWN
            ),
            diagnostics=layout.diagnostics,
            reporting_period_start=layout.reporting_period_start,
            reporting_period_end=layout.reporting_period_end,
        )

    def parse(self, document: ImportDocument) -> ParsedImport:
        layout = _read_layout(document)
        if not layout.matched:
            return ParsedImport(
                rows=(
                    ParsedRow(
                        sequence_number=1,
                        raw_data={},
                        errors=layout.diagnostics
                        or (
                            {
                                "code": "tbank_xlsx_signature_invalid",
                                "message": "Workbook does not match T-Investments XLSX v1",
                            },
                        ),
                    ),
                )
            )

        instrument_isins = dict(layout.instrument_isins)
        instrument_metadata = dict(layout.instrument_metadata)
        parsed_rows: list[ParsedRow] = []
        sequence_number = 1
        for source_row in layout.trade_rows:
            rows = _parse_trade_row(
                source_row,
                sequence_number,
                instrument_isins,
                instrument_metadata,
            )
            parsed_rows.extend(rows)
            sequence_number += len(rows)
        for source_row in layout.cash_movement_rows:
            parsed_rows.append(_parse_cash_movement(source_row, sequence_number))
            sequence_number += 1
        for source_row in layout.cash_rows:
            parsed_rows.append(_parse_cash_control(source_row, sequence_number))
            sequence_number += 1
        for source_row in layout.security_rows:
            parsed_rows.append(_parse_security_control(source_row, sequence_number))
            sequence_number += 1
        return ParsedImport(rows=tuple(parsed_rows), diagnostics=layout.diagnostics)

    def validate(self, parsed: ParsedImport) -> ValidationResult:
        return ValidationResult(rows=parsed.rows, diagnostics=parsed.diagnostics)


def _read_layout(document: ImportDocument) -> _WorkbookLayout:
    try:
        _validate_zip_container(document)
        document.stream.seek(0)
        workbook = load_workbook(document.stream, read_only=True, data_only=True)
    except (BadZipFile, KeyError, OSError, ValueError):
        document.stream.seek(0)
        return _layout_error(
            "tbank_xlsx_container_invalid",
            "Workbook is not a valid bounded XLSX document",
        )

    try:
        if "broker_rep" not in workbook.sheetnames:
            if _MULTISHEET_REQUIRED_SHEETS.issubset(workbook.sheetnames):
                return _read_multisheet_layout(workbook, document.original_filename)
            return _layout_error(
                "tbank_xlsx_sheet_missing",
                "Workbook does not contain the required broker report sheet",
            )
        sheet = workbook["broker_rep"]
        sheet.calculate_dimension(force=True)
        if (sheet.max_row or 0) > 100_000 or (sheet.max_column or 0) > 512:
            return _layout_error(
                "tbank_xlsx_sheet_limit_exceeded",
                "Broker report sheet exceeds configured structural limits",
            )
        rows = tuple(tuple(row) for row in sheet.iter_rows(values_only=True))
    finally:
        workbook.close()
        document.stream.seek(0)

    title_match: re.Match[str] | None = None
    for row in rows:
        first = _first_text(row)
        if first is not None and (match := _REPORT_TITLE.fullmatch(first)) is not None:
            title_match = match
            break
    if title_match is None:
        return _layout_error(
            "tbank_xlsx_title_missing",
            "Workbook does not contain the required broker report title",
        )

    required_sections = {
        "trades": "1.1 Информация о совершенных и исполненных сделках",
        "cash": "2. Операции с денежными средствами",
        "securities": "3.1 Движение по ценным бумагам инвестора",
    }
    section_rows: dict[str, int] = {}
    for key, prefix in required_sections.items():
        row_index = _find_row_by_prefix(rows, prefix)
        if row_index is None:
            return _layout_error(
                "tbank_xlsx_section_missing",
                "Workbook is missing a required T-Investments report section",
            )
        section_rows[key] = row_index

    trade_header_index = _find_header_row(
        rows,
        start=section_rows["trades"] + 1,
        required=_REQUIRED_TRADE_HEADERS,
    )
    if trade_header_index is None:
        return _layout_error(
            "tbank_xlsx_trade_headers_missing",
            "Executed trades section is missing required columns",
        )
    header_map = _header_map(rows[trade_header_index])
    trade_rows = _read_data_rows(rows, trade_header_index + 1, header_map)

    cash_rows = _section_data_rows(
        rows,
        section_rows["cash"],
        frozenset({"Валюта", "Входящий остаток", "Исходящий остаток"}),
    )
    security_rows = _section_data_rows(
        rows,
        section_rows["securities"],
        frozenset(
            {
                "Наименование актива",
                "Код актива",
                "ISIN",
                "Входящий остаток",
                "Зачисление",
                "Списание",
                "Исходящий остаток",
            }
        ),
    )

    instrument_isins: dict[str, str] = {}
    instrument_section = _find_row_by_prefix(rows, "4.1 Информация о ценных бумагах")
    if instrument_section is not None:
        instrument_header = _find_header_row(
            rows,
            start=instrument_section + 1,
            required=frozenset({"Код актива", "ISIN"}),
        )
        if instrument_header is not None:
            for item in _read_data_rows(
                rows,
                instrument_header + 1,
                _header_map(rows[instrument_header]),
            ):
                code = _text(item.values.get("Код актива"))
                isin = _text(item.values.get("ISIN")).upper()
                if code and isin:
                    instrument_isins[code.upper()] = isin

    diagnostics: list[Diagnostic] = [
        {
            "code": "tbank_cash_reconciliation_available",
            "message": "Cash balances were read as reconciliation data only",
        },
        {
            "code": "tbank_securities_reconciliation_available",
            "message": "Security movements were read as reconciliation data only",
        },
    ]
    for prefix, code in (
        ("1.2 Информация о неисполненных сделках", "tbank_unfinished_trades_empty"),
        ("1.3 Сделки за расчетный период", "tbank_terminated_trades_empty"),
        ("3.2 Движение по производным", "tbank_derivatives_empty"),
    ):
        section = _find_row_by_prefix(rows, prefix)
        if section is not None and _next_nonempty_is_empty_marker(rows, section + 1):
            diagnostics.append(
                {"code": code, "message": "Optional report section contains no data"}
            )
    unsupported_section_count = sum(
        1
        for row in rows
        if (value := _first_text(row)) is not None
        and _is_section_title(value)
        and not value.startswith(_KNOWN_SECTION_PREFIXES)
    )
    if unsupported_section_count:
        diagnostics.append(
            {
                "code": "tbank_section_unsupported",
                "message": "Report contains sections not normalized by adapter version 1.0",
                "count": unsupported_section_count,
            }
        )

    return _WorkbookLayout(
        matched=True,
        reporting_period_start=datetime.strptime(
            title_match.group("start"), "%d.%m.%Y"
        ).date(),
        reporting_period_end=datetime.strptime(title_match.group("end"), "%d.%m.%Y").date(),
        trade_rows=trade_rows,
        cash_rows=cash_rows,
        security_rows=security_rows,
        instrument_isins=tuple(sorted(instrument_isins.items())),
        diagnostics=tuple(diagnostics),
    )


def _read_multisheet_layout(workbook: object, original_filename: str) -> _WorkbookLayout:
    try:
        sheets = {
            name: tuple(
                tuple(row)
                for row in workbook[name].iter_rows(values_only=True)  # type: ignore[index]
            )
            for name in _MULTISHEET_REQUIRED_SHEETS
        }
    except (KeyError, OSError, ValueError):
        return _layout_error(
            "tbank_multisheet_structure_invalid",
            "T-Investments workbook contains an invalid multi-sheet structure",
        )
    if any(
        len(rows) > 100_000
        or max((len(row) for row in rows), default=0) > 512
        for rows in sheets.values()
    ):
        return _layout_error(
            "tbank_xlsx_sheet_limit_exceeded",
            "Broker report sheet exceeds configured structural limits",
        )

    period = _multisheet_period(original_filename)
    if period is None:
        return _layout_error(
            "tbank_xlsx_period_missing",
            "T-Investments workbook filename does not contain the reporting period",
        )

    trade_rows_source = sheets["Завершенные сделки"]
    trade_header = _find_header_anywhere(
        trade_rows_source,
        frozenset(
            {
                "№ сделки",
                "Дата заключен.",
                "ISIN/рег.код",
                "Актив",
                "Количество актива⁷, шт./грамм",
                "Цена",
                "Валюта расчетов",
                "Комиссия банка",
            }
        ),
    )
    if trade_header is None:
        return _layout_error(
            "tbank_xlsx_trade_headers_missing",
            "Executed trades sheet is missing required columns",
        )
    trade_columns = _header_map(trade_rows_source[trade_header])

    position_rows_source = sheets["Динамика позиций"]
    position_header = _find_header_anywhere(
        position_rows_source,
        frozenset({"Инструмент", "Актив ¹", "кол-во, шт./грамм⁵"}),
    )
    position_types_by_name: dict[str, str] = {}
    position_values_by_name: dict[str, tuple[int, object, object]] = {}
    if position_header is not None:
        current_type: str | None = None
        for row_index, row in enumerate(
            position_rows_source[position_header + 2 :],
            start=position_header + 3,
        ):
            group = _cell(row, 2)
            if (mapped := _tbank_group_type(_text(group))) is not None:
                current_type = mapped
            name = _text(_cell(row, 6))
            opening = _cell(row, 12)
            closing = _cell(row, 15)
            if not name or current_type is None or (opening is None and closing is None):
                continue
            position_types_by_name[name] = current_type
            position_values_by_name[name] = (row_index, opening, closing)

    trade_rows: list[_SourceRow] = []
    instrument_isins: dict[str, str] = {}
    instrument_metadata: dict[str, _InstrumentMetadata] = {}
    code_by_name: dict[str, str] = {}
    for row_index, row in enumerate(
        trade_rows_source[trade_header + 1 :],
        start=trade_header + 2,
    ):
        trade_number = _text(_column(row, trade_columns, "№ сделки"))
        if not trade_number:
            if trade_rows:
                break
            continue
        raw_quantity = _column(row, trade_columns, "Количество актива⁷, шт./грамм")
        quantity = _signed_source_decimal(raw_quantity)
        raw_datetime = _text(_column(row, trade_columns, "Дата заключен."))
        date_value, time_value = _split_multisheet_datetime(raw_datetime)
        code = _text(_column(row, trade_columns, "ISIN/рег.код")).upper()
        name = _text(_column(row, trade_columns, "Актив"))
        settlement_currency = _text(
            _column(row, trade_columns, "Валюта расчетов")
        ).upper()
        fee = _column(row, trade_columns, "Комиссия банка")
        fee_currency = _text(
            _column(row, trade_columns, "Валюта комиссии")
        ).upper()
        values = {
            header: _column(row, trade_columns, header) for header in trade_columns
        }
        values.update(
            {
                "Номер сделки": trade_number,
                "Дата заключения": date_value,
                "Время": time_value,
                "Вид сделки": (
                    "Покупка" if quantity is not None and quantity > 0 else "Продажа"
                ),
                "Наименование актива": name,
                "Код актива": code,
                "Цена за единицу": _column(row, trade_columns, "Цена"),
                "Валюта цены": settlement_currency,
                "Количество": raw_quantity,
                "Комиссия брокера": fee,
                "Валюта комиссии": fee_currency or settlement_currency,
            }
        )
        trade_rows.append(
            _SourceRow(
                source_row_number=row_index,
                values=values,
                source_sheet="Завершенные сделки",
            )
        )
        if _ISIN.fullmatch(code):
            instrument_isins[code] = code
        if code and name:
            code_by_name[name] = code
            instrument_metadata[code] = _InstrumentMetadata(
                name=name[:200],
                instrument_type=position_types_by_name.get(name, _infer_tbank_type(name)),
            )

    cash_sheet = sheets[" Движение ДС"]
    cash_header = _find_header_anywhere(
        cash_sheet,
        frozenset({"Дата", "Наименование операции", "Комментарий"}),
    )
    cash_movement_rows: list[_SourceRow] = []
    cash_controls: list[_SourceRow] = []
    if cash_header is not None:
        for row_index, row in enumerate(
            cash_sheet[cash_header + 2 :],
            start=cash_header + 3,
        ):
            operation_name = _text(_cell(row, 9))
            if operation_name != "Перевод":
                continue
            occurred_at = _cell(row, 6) or _cell(row, 2)
            amount = _cell(row, 14)
            if occurred_at is None or amount is None:
                continue
            cash_movement_rows.append(
                _SourceRow(
                    source_row_number=row_index,
                    source_sheet=" Движение ДС",
                    values={
                        "Дата": occurred_at,
                        "Наименование операции": operation_name,
                        "Комментарий": _cell(row, 10),
                        "Сумма": amount,
                        "Валюта": "RUB",
                    },
                )
            )
        opening = _cell(cash_sheet[cash_header - 2], 14) if cash_header >= 2 else None
        closing = next(
            (
                _cell(row, 17)
                for row in reversed(cash_sheet[cash_header + 2 :])
                if isinstance(_cell(row, 17), int | float | Decimal)
            ),
            None,
        )
        if opening is not None and closing is not None:
            cash_controls.append(
                _SourceRow(
                    source_row_number=cash_header + 1,
                    source_sheet=" Движение ДС",
                    values={
                        "Валюта": "RUB",
                        "Входящий остаток": opening,
                        "Исходящий остаток": closing,
                    },
                )
            )

    security_rows: list[_SourceRow] = []
    for name, (row_index, opening, closing) in position_values_by_name.items():
        position_code = code_by_name.get(name)
        if position_code is None:
            continue
        security_rows.append(
            _SourceRow(
                source_row_number=row_index,
                source_sheet="Динамика позиций",
                values={
                    "Наименование актива": name,
                    "Код актива": position_code,
                    "ISIN": position_code if _ISIN.fullmatch(position_code) else "",
                    "Входящий остаток": opening,
                    "Зачисление": None,
                    "Списание": None,
                    "Исходящий остаток": closing,
                },
            )
        )

    return _WorkbookLayout(
        matched=True,
        reporting_period_start=period[0],
        reporting_period_end=period[1],
        trade_rows=tuple(trade_rows),
        cash_movement_rows=tuple(cash_movement_rows),
        cash_rows=tuple(cash_controls),
        security_rows=tuple(security_rows),
        instrument_isins=tuple(sorted(instrument_isins.items())),
        instrument_metadata=tuple(sorted(instrument_metadata.items())),
        diagnostics=(
            {
                "code": "tbank_multisheet_layout_detected",
                "message": "Current T-Investments multi-sheet XLSX layout was detected",
            },
            {
                "code": "tbank_cash_reconciliation_available",
                "message": "Cash balances were read as reconciliation data only",
            },
            {
                "code": "tbank_securities_reconciliation_available",
                "message": "Security movements were read as reconciliation data only",
            },
        ),
    )


def _find_header_anywhere(
    rows: tuple[tuple[object, ...], ...], required: frozenset[str]
) -> int | None:
    for index, row in enumerate(rows):
        if required.issubset(_header_map(row)):
            return index
    return None


def _cell(row: tuple[object, ...], zero_based_index: int) -> object | None:
    return row[zero_based_index] if zero_based_index < len(row) else None


def _column(
    row: tuple[object, ...], columns: dict[str, int], name: str
) -> object | None:
    index = columns.get(name)
    return _cell(row, index) if index is not None else None


def _multisheet_period(filename: str) -> tuple[date, date] | None:
    match = _MULTISHEET_PERIOD.search(filename)
    if match is None:
        return None
    try:
        start = _short_date(match.group("start"))
        end = _short_date(match.group("end"))
    except ValueError:
        return None
    return (start, end) if start <= end else None


def _short_date(value: str) -> date:
    return datetime.strptime(value, "%d.%m.%y" if len(value) == 8 else "%d.%m.%Y").date()


def _split_multisheet_datetime(value: str) -> tuple[str, str]:
    normalized = " ".join(value.split())
    try:
        parsed = datetime.strptime(normalized, "%d.%m.%Y %H:%M:%S")
    except ValueError:
        return value, ""
    return parsed.strftime("%d.%m.%Y"), parsed.strftime("%H:%M:%S")


def _signed_source_decimal(value: object) -> Decimal | None:
    text = _text(value).replace("\u00a0", "").replace(" ", "").replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _tbank_group_type(value: str) -> str | None:
    return {
        "Акции": "stock",
        "Облигации": "bond",
        "Прочее": "fund",
    }.get(value)


def _infer_tbank_type(name: str) -> str:
    normalized = name.casefold()
    if "облигац" in normalized or "офз" in normalized:
        return "bond"
    if "etf" in normalized:
        return "etf"
    if any(
        marker in normalized
        for marker in ("бпиф", "опиф", "пиф", "фонд", "денежн", "накопительн")
    ):
        return "fund"
    if "опцион" in normalized:
        return "option"
    return "stock"


def _validate_zip_container(document: ImportDocument) -> None:
    document.stream.seek(0)
    with ZipFile(document.stream) as archive:
        members = archive.infolist()
        names = {member.filename for member in members}
        if len(members) > _MAX_ARCHIVE_ENTRIES:
            raise ValueError
        if sum(member.file_size for member in members) > _MAX_UNCOMPRESSED_BYTES:
            raise ValueError
        if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
            raise ValueError
    document.stream.seek(0)


def _parse_trade_row(
    source: _SourceRow,
    sequence_number: int,
    instrument_isins: dict[str, str],
    instrument_metadata: dict[str, _InstrumentMetadata],
) -> tuple[ParsedRow, ...]:
    errors: list[Diagnostic] = []
    warnings: list[Diagnostic] = []
    trade_number = _text(source.values.get("Номер сделки"))
    if not trade_number or len(trade_number) > 256:
        errors.append({"code": "trade_number_invalid", "message": "Trade number is required"})
    side_value = _text(source.values.get("Вид сделки")).casefold()
    side = {"покупка": "buy", "продажа": "sell"}.get(side_value)
    if side is None:
        errors.append(
            {"code": "trade_side_unsupported", "message": "Trade side is not supported"}
        )
    occurred_at = _parse_occurred_at(
        source.values.get("Дата заключения"),
        source.values.get("Время"),
        errors,
    )
    quantity = _required_decimal(
        source.values.get("Количество"), _QUANTITY_ADAPTER, "Quantity", errors
    )
    price = _required_decimal(
        source.values.get("Цена за единицу"), _PRICE_ADAPTER, "Price", errors
    )
    currency = _currency(source.values.get("Валюта цены"), errors)
    code = _text(source.values.get("Код актива")).upper()
    if not code:
        errors.append(
            {"code": "instrument_provider_code_missing", "message": "Asset code is required"}
        )
    reference: dict[str, object] = {
        "provider": "tbank_broker_xlsx",
        "provider_code": code,
        "currency": currency,
    }
    isin = instrument_isins.get(code)
    if isin is not None and _ISIN.fullmatch(isin):
        reference["isin"] = isin
    metadata = instrument_metadata.get(code)
    if metadata is not None:
        reference.update(
            {
                "auto_create": True,
                "name": metadata.name,
                "instrument_type": metadata.instrument_type,
            }
        )

    raw_data = {key: _raw_value(value) for key, value in source.values.items()}
    if errors or side is None or occurred_at is None or quantity is None or price is None:
        return (
            ParsedRow(
                sequence_number=sequence_number,
                source_sheet=source.source_sheet,
                source_row_number=source.source_row_number,
                raw_data=raw_data,
                warnings=tuple(warnings),
                errors=tuple(errors),
            ),
        )

    common: dict[str, object] = {
        "occurred_at": occurred_at,
        "time_precision": "second",
        "source_operation_id": trade_number,
    }
    candidates: list[dict[str, object]] = [
        {
            **common,
            "operation_type": "trade",
            "payload": {
                "side": side,
                "quantity": quantity,
                "price": price,
                "price_currency": currency,
                "instrument_reference": reference,
            },
        }
    ]
    fees_by_currency: dict[str, Decimal] = {}
    for amount_field, currency_field in _FEE_FIELDS:
        fee = _optional_decimal(
            source.values.get(amount_field),
            _MONEY_ADAPTER,
            amount_field,
            errors,
        )
        if fee is None or fee == 0:
            continue
        fee_currency = _currency(source.values.get(currency_field), errors)
        if fee_currency is not None:
            fees_by_currency[fee_currency] = (
                fees_by_currency.get(fee_currency, Decimal(0)) + abs(fee)
            )
    for fee_currency, fee in sorted(fees_by_currency.items()):
        candidates.append(
            {
                **common,
                "source_operation_id": f"{trade_number}:fee:{fee_currency}",
                "operation_type": "fee",
                "payload": {
                    "amount": decimal_to_json(fee),
                    "currency": fee_currency,
                    "instrument_reference": reference,
                },
            }
        )
    if errors:
        return (
            ParsedRow(
                sequence_number=sequence_number,
                source_sheet=source.source_sheet,
                source_row_number=source.source_row_number,
                raw_data=raw_data,
                warnings=tuple(warnings),
                errors=tuple(errors),
            ),
        )
    return tuple(
        ParsedRow(
            sequence_number=sequence_number + offset,
            source_sheet=source.source_sheet,
            source_row_number=source.source_row_number,
            raw_data=raw_data,
            normalized_candidate=candidate,
            reconciliation_data=(
                {
                    "position_effects": [
                        {
                            "key": _instrument_key(reference),
                            "change": (
                                decimal_to_json(Decimal(quantity))
                                if side == "buy"
                                else decimal_to_json(-Decimal(quantity))
                            ),
                        }
                    ]
                }
                if offset == 0
                else None
            ),
            warnings=tuple(warnings if offset == 0 else ()),
        )
        for offset, candidate in enumerate(candidates)
    )


def _parse_cash_control(source: _SourceRow, sequence_number: int) -> ParsedRow:
    errors: list[Diagnostic] = []
    currency = _currency(source.values.get("Валюта"), errors)
    opening = _control_decimal(
        source.values.get("Входящий остаток"), _MONEY_ADAPTER, "Opening cash", errors
    )
    closing = _control_decimal(
        source.values.get("Исходящий остаток"), _MONEY_ADAPTER, "Closing cash", errors
    )
    return _control_row(
        source,
        sequence_number,
        code="tbank_cash_reconciliation_only",
        message="Cash balance is reconciliation data only",
        errors=errors,
        control=(
            None
            if currency is None or opening is None or closing is None
            else {
                "kind": "cash",
                "key": currency,
                "opening": decimal_to_json(opening),
                "closing": decimal_to_json(closing),
            }
        ),
    )


def _parse_cash_movement(source: _SourceRow, sequence_number: int) -> ParsedRow:
    errors: list[Diagnostic] = []
    raw_occurred_at = source.values.get("Дата")
    if isinstance(raw_occurred_at, datetime):
        occurred_at = raw_occurred_at.replace(tzinfo=_TIMEZONE).isoformat()
        precision = "second"
    elif isinstance(raw_occurred_at, date):
        occurred_at = datetime.combine(
            raw_occurred_at, time.min, tzinfo=_TIMEZONE
        ).isoformat()
        precision = "date"
    else:
        text = " ".join(_text(raw_occurred_at).split())
        occurred_at = None
        precision = "second"
        for pattern, candidate_precision in (
            ("%d.%m.%Y %H:%M:%S", "second"),
            ("%d.%m.%Y", "date"),
        ):
            try:
                occurred_at = datetime.strptime(text, pattern).replace(
                    tzinfo=_TIMEZONE
                ).isoformat()
                precision = candidate_precision
                break
            except ValueError:
                continue
        if occurred_at is None:
            errors.append(
                {
                    "code": "tbank_cash_datetime_invalid",
                    "message": "Cash movement date is invalid",
                }
            )
    amount = _signed_source_decimal(source.values.get("Сумма"))
    if amount is None or amount == 0:
        errors.append(
            {
                "code": "tbank_cash_amount_invalid",
                "message": "Cash movement amount must be a non-zero decimal",
            }
        )
    currency = _currency(source.values.get("Валюта"), errors)
    raw_data = {key: _raw_value(value) for key, value in source.values.items()}
    if errors or occurred_at is None or amount is None or amount == 0 or currency is None:
        return ParsedRow(
            sequence_number=sequence_number,
            source_sheet=source.source_sheet,
            source_row_number=source.source_row_number,
            raw_data=raw_data,
            errors=tuple(errors),
        )
    return ParsedRow(
        sequence_number=sequence_number,
        source_sheet=source.source_sheet,
        source_row_number=source.source_row_number,
        raw_data=raw_data,
        normalized_candidate={
            "occurred_at": occurred_at,
            "time_precision": precision,
            "operation_type": "cash_movement",
            "payload": {
                "direction": "deposit" if amount > 0 else "withdrawal",
                "amount": decimal_to_json(abs(amount)),
                "currency": currency,
            },
        },
    )


def _parse_security_control(source: _SourceRow, sequence_number: int) -> ParsedRow:
    errors: list[Diagnostic] = []
    code = _text(source.values.get("Код актива")).upper()
    raw_isin = _text(source.values.get("ISIN")).upper()
    if not code and _ISIN.fullmatch(raw_isin) is None:
        errors.append(
            {
                "code": "instrument_provider_code_missing",
                "message": "Security control has no supported instrument identifier",
            }
        )
    opening = _control_decimal(
        source.values.get("Входящий остаток"),
        _QUANTITY_ADAPTER,
        "Opening position",
        errors,
    )
    closing = _control_decimal(
        source.values.get("Исходящий остаток"),
        _QUANTITY_ADAPTER,
        "Closing position",
        errors,
    )
    instrument_key = (
        f"isin:{raw_isin}" if _ISIN.fullmatch(raw_isin) else f"provider_code:{code}"
    )
    return _control_row(
        source,
        sequence_number,
        code="tbank_security_reconciliation_only",
        message="Security movement is reconciliation data only",
        errors=errors,
        control=(
            None
            if errors or opening is None or closing is None
            else {
                "kind": "position",
                "key": instrument_key,
                "opening": decimal_to_json(opening),
                "closing": decimal_to_json(closing),
            }
        ),
    )


def _control_row(
    source: _SourceRow,
    sequence_number: int,
    *,
    code: str,
    message: str,
    errors: list[Diagnostic],
    control: dict[str, object] | None,
) -> ParsedRow:
    return ParsedRow(
        sequence_number=sequence_number,
        source_sheet=source.source_sheet,
        source_row_number=source.source_row_number,
        raw_data={key: _raw_value(value) for key, value in source.values.items()},
        reconciliation_data={"controls": [control]} if control is not None else None,
        status=ImportRowStatus.ERROR if errors else ImportRowStatus.EXCLUDED,
        warnings=() if errors else ({"code": code, "message": message},),
        errors=tuple(errors),
    )


def _instrument_key(reference: dict[str, object]) -> str:
    isin = reference.get("isin")
    if isinstance(isin, str):
        return f"isin:{isin}"
    return f"provider_code:{reference['provider_code']}"


def _read_data_rows(
    rows: tuple[tuple[object, ...], ...],
    start: int,
    header_map: dict[str, int],
) -> tuple[_SourceRow, ...]:
    result: list[_SourceRow] = []
    for row_index in range(start, len(rows)):
        first = _first_text(rows[row_index])
        if first is None:
            if result:
                break
            continue
        if _is_section_title(first) or first.casefold().startswith("нет данных"):
            break
        values = {
            header: rows[row_index][column] if column < len(rows[row_index]) else None
            for header, column in header_map.items()
        }
        result.append(_SourceRow(source_row_number=row_index + 1, values=values))
    return tuple(result)


def _section_data_rows(
    rows: tuple[tuple[object, ...], ...],
    section_row: int,
    required_headers: frozenset[str],
) -> tuple[_SourceRow, ...]:
    header_row = _find_header_row(
        rows,
        start=section_row + 1,
        required=required_headers,
    )
    if header_row is None:
        return ()
    return _read_data_rows(rows, header_row + 1, _header_map(rows[header_row]))


def _find_header_row(
    rows: tuple[tuple[object, ...], ...],
    *,
    start: int,
    required: frozenset[str],
) -> int | None:
    for row_index in range(start, min(start + 6, len(rows))):
        if required.issubset(_header_map(rows[row_index])):
            return row_index
    return None


def _header_map(row: tuple[object, ...]) -> dict[str, int]:
    return {
        " ".join(value.split()): index
        for index, item in enumerate(row)
        if (value := _text(item))
    }


def _find_row_by_prefix(rows: tuple[tuple[object, ...], ...], prefix: str) -> int | None:
    for index, row in enumerate(rows):
        value = _first_text(row)
        if value is not None and value.startswith(prefix):
            return index
    return None


def _next_nonempty_is_empty_marker(
    rows: tuple[tuple[object, ...], ...], start: int
) -> bool:
    for row in rows[start : start + 4]:
        value = _first_text(row)
        if value is not None and value.casefold().startswith("нет данных"):
            return True
    return False


def _is_section_title(value: str) -> bool:
    return re.match(r"^[1-9]\d*(?:\.\d+)?[ .]", value) is not None


def _first_text(row: tuple[object, ...]) -> str | None:
    for item in row:
        value = _text(item)
        if value:
            return value
    return None


def _parse_occurred_at(
    raw_date: object,
    raw_time: object,
    errors: list[Diagnostic],
) -> str | None:
    try:
        parsed_date = (
            raw_date.date()
            if isinstance(raw_date, datetime)
            else raw_date
            if isinstance(raw_date, date)
            else datetime.strptime(_text(raw_date), "%d.%m.%Y").date()
        )
        parsed_time = (
            raw_time.time()
            if isinstance(raw_time, datetime)
            else raw_time
            if isinstance(raw_time, time)
            else datetime.strptime(_text(raw_time), "%H:%M:%S").time()
        )
    except ValueError:
        errors.append(
            {"code": "trade_datetime_invalid", "message": "Trade date or time is invalid"}
        )
        return None
    return datetime.combine(parsed_date, parsed_time, tzinfo=_TIMEZONE).isoformat()


def _currency(value: object, errors: list[Diagnostic]) -> str | None:
    try:
        return _CURRENCY_ADAPTER.validate_python(_text(value).upper())
    except ValidationError:
        errors.append({"code": "currency_invalid", "message": "Currency is invalid"})
        return None


def _required_decimal(
    value: object,
    adapter: TypeAdapter[Decimal],
    field_name: str,
    errors: list[Diagnostic],
) -> str | None:
    parsed = _optional_decimal(value, adapter, field_name, errors)
    if parsed is None:
        return None
    if parsed == 0:
        errors.append(
            {"code": "value_not_positive", "message": f"{field_name} must not be zero"}
        )
        return None
    return decimal_to_json(abs(parsed))


def _optional_decimal(
    value: object,
    adapter: TypeAdapter[Decimal],
    field_name: str,
    errors: list[Diagnostic],
) -> Decimal | None:
    text = _text(value).replace("\u00a0", "").replace(" ", "").replace(",", ".")
    if not text:
        return None
    try:
        parsed = Decimal(text)
        return adapter.validate_python(decimal_to_json(abs(parsed)))
    except (InvalidOperation, ValidationError):
        errors.append(
            {"code": "decimal_value_invalid", "message": f"{field_name} is invalid"}
        )
        return None


def _control_decimal(
    value: object,
    adapter: TypeAdapter[Decimal],
    field_name: str,
    errors: list[Diagnostic],
) -> Decimal | None:
    text = _text(value).replace("\u00a0", "").replace(" ", "").replace(",", ".")
    if not text:
        errors.append(
            {"code": "decimal_value_missing", "message": f"{field_name} is required"}
        )
        return None
    try:
        parsed = Decimal(text)
        return adapter.validate_python(decimal_to_json(parsed))
    except (InvalidOperation, ValidationError):
        errors.append(
            {"code": "decimal_value_invalid", "message": f"{field_name} is invalid"}
        )
        return None


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return repr(value)
    result = str(value).strip()
    return result[1:].strip() if result.startswith("'") else result


def _raw_value(value: object) -> object:
    if value is None or isinstance(value, str | int | bool):
        return value
    if isinstance(value, Decimal):
        return decimal_to_json(value)
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    return _text(value)


def _layout_error(code: str, message: str) -> _WorkbookLayout:
    return _WorkbookLayout(matched=False, diagnostics=({"code": code, "message": message},))
